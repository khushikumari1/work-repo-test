"""
logs/cost_logger.py — Tracks per-run token usage and USD cost for the
Test Bank Audit Pipeline, and appends each run as a row in an Excel
workbook (logs/cost_log.xlsx).

Pricing lives in MODEL_PRICING below. Update it whenever Google changes
Gemini pricing, or add a new entry when you switch models — an unrecognised
model_id logs a warning and is costed at $0 rather than raising, so a
pricing gap never crashes a pipeline run.
"""

import logging
import time
import uuid
from dataclasses import dataclass, fields
from datetime import datetime
from pathlib import Path
from typing import Dict, Tuple

from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)

# =====================================================================
# Pricing table — USD per 1,000,000 tokens: (input_rate, output_rate)
# =====================================================================
MODEL_PRICING: Dict[str, Tuple[float, float]] = {
    "gemini-3.5-flash": (1.50, 9.00),
    "gemini-3.5-flash-lite": (0.30, 2.50),
}
DEFAULT_PRICING: Tuple[float, float] = (0.0, 0.0)

SHEET_NAME = "Runs"


def get_pricing(model_id: str) -> Tuple[float, float]:
    """Look up (input_rate, output_rate) per 1M tokens for a model_id.
    Falls back to $0/$0 (with a logged warning) for unrecognised models,
    so a missing pricing entry never crashes a run — it just shows up as
    an obviously-wrong $0.00 line in the log for you to notice and fix."""
    key = (model_id or "").strip().lower()
    if key not in MODEL_PRICING:
        logger.warning(
            "No pricing entry for model '%s' — cost will be logged as $0.00. "
            "Add an entry to MODEL_PRICING in logs/cost_logger.py.", model_id,
        )
        return DEFAULT_PRICING
    return MODEL_PRICING[key]


@dataclass
class RunCostRecord:
    """One row of the cost log — a full summary of a single pipeline run."""

    run_id: str
    timestamp: str
    model_id: str
    region: str
    test_bank_file: str
    textbook_file: str
    duration_s: float
    num_chunks: int
    retried_chunks: int
    total_input_tokens: int
    total_output_tokens: int
    total_tokens: int
    input_cost: float
    output_cost: float
    total_cost: float
    total_items: int
    keep_count: int
    revise_count: int
    remove_count: int
    parse_error_count: int


class CostAccumulator:
    """Accumulates per-chunk token usage over the life of one pipeline run,
    then produces a single RunCostRecord via finalise()."""

    def __init__(self, model_id: str, region: str, test_bank_file: str, textbook_file: str):
        self.model_id = model_id
        self.region = region
        self.test_bank_file = str(test_bank_file)
        self.textbook_file = str(textbook_file)
        self.run_id = uuid.uuid4().hex[:12]
        self._start = time.time()
        self._chunks = []  # objects with .input_tokens / .output_tokens / .retried

    def add_chunk(self, chunk_stats) -> None:
        """Record one chunk's usage. Accepts anything with input_tokens,
        output_tokens, and retried attributes (e.g. testbank_pipeline's
        ChunkStats) — duck-typed on purpose so this module never needs to
        import testbank_pipeline and create a circular import."""
        self._chunks.append(chunk_stats)

    def finalise(self, total_items: int, keep: int, revise: int, remove: int, parse_errors: int) -> RunCostRecord:
        total_input = sum(getattr(c, "input_tokens", 0) for c in self._chunks)
        total_output = sum(getattr(c, "output_tokens", 0) for c in self._chunks)
        retried = sum(1 for c in self._chunks if getattr(c, "retried", False))

        input_rate, output_rate = get_pricing(self.model_id)
        input_cost = total_input / 1_000_000 * input_rate
        output_cost = total_output / 1_000_000 * output_rate

        return RunCostRecord(
            run_id=self.run_id,
            timestamp=datetime.now().isoformat(timespec="seconds"),
            model_id=self.model_id,
            region=self.region,
            test_bank_file=self.test_bank_file,
            textbook_file=self.textbook_file,
            duration_s=round(time.time() - self._start, 2),
            num_chunks=len(self._chunks),
            retried_chunks=retried,
            total_input_tokens=total_input,
            total_output_tokens=total_output,
            total_tokens=total_input + total_output,
            input_cost=round(input_cost, 6),
            output_cost=round(output_cost, 6),
            total_cost=round(input_cost + output_cost, 6),
            total_items=total_items,
            keep_count=keep,
            revise_count=revise,
            remove_count=remove,
            parse_error_count=parse_errors,
        )


def append_to_excel(record: RunCostRecord, log_path: str) -> None:
    """Append one RunCostRecord as a row to logs/cost_log.xlsx, creating
    the workbook and header row on first use."""
    path = Path(log_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = [f.name for f in fields(RunCostRecord)]

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.create_sheet(SHEET_NAME)
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            sheet.append(headers)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(headers)

    sheet.append([getattr(record, h) for h in headers])
    workbook.save(path)
    logger.info("Appended run %s to %s", record.run_id, path)