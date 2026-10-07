"""
logs/output_logger.py — Appends one row per test-bank item decision
(KEEP / REVISE / REMOVE / PARSE_ERROR) to logs/output_log.xlsx for every
pipeline run, so a human editor can audit what changed without opening
the output .docx.
"""

import logging
from pathlib import Path
from typing import List

from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)

SHEET_NAME = "Decisions"
HEADERS = [
    "run_id", "timestamp", "test_bank_file", "duration_s",
    "item_number", "decision", "flags",
]


def append_output_log(
    run_id: str,
    timestamp: str,
    test_bank_file: str,
    output_items: List,
    duration_s: float,
    log_path: str,
) -> None:
    """Append one row per OutputItem to logs/output_log.xlsx, creating the
    workbook and header row on first use. `output_items` is the list of
    OutputItem objects from testbank_pipeline — only .number, .decision,
    and .flags are read, so this module stays decoupled from that dataclass
    and doesn't need to import it."""
    path = Path(log_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.create_sheet(SHEET_NAME)
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            sheet.append(HEADERS)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)

    for item in output_items:
        flags = getattr(item, "flags", None) or []
        sheet.append([
            run_id,
            timestamp,
            str(test_bank_file),
            duration_s,
            getattr(item, "number", ""),
            getattr(item, "decision", ""),
            "; ".join(flags),
        ])

    workbook.save(path)
    logger.info("Logged %d item decisions for run %s to %s", len(output_items), run_id, path)