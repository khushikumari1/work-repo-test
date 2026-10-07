"""
logs/calculate_logs.py — Rebuilds the "Summary" sheet in logs/cost_log.xlsx
from the raw "Runs" sheet written by cost_logger.append_to_excel().

Run standalone with:
    python logs/calculate_logs.py logs/cost_log.xlsx
"""

import logging
import sys
from pathlib import Path
from typing import Dict

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

RUNS_SHEET = "Runs"
SUMMARY_SHEET = "Summary"


def calculate(cost_log_path: str) -> str:
    """Read the Runs sheet of the cost log, recompute totals/averages and
    a per-model breakdown, and (re)write them into a Summary sheet in the
    same workbook. Returns the path it wrote to."""
    path = Path(cost_log_path)
    if not path.exists():
        raise FileNotFoundError(f"Cost log not found: {cost_log_path}")

    workbook = load_workbook(path)
    if RUNS_SHEET not in workbook.sheetnames:
        raise ValueError(f"'{RUNS_SHEET}' sheet not found in {cost_log_path}")

    runs_sheet = workbook[RUNS_SHEET]
    header_row = next(runs_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [h for h in header_row if h is not None]

    rows = []
    for row in runs_sheet.iter_rows(min_row=2, values_only=True):
        if row and any(v is not None for v in row):
            rows.append(dict(zip(headers, row)))

    if SUMMARY_SHEET in workbook.sheetnames:
        del workbook[SUMMARY_SHEET]
    summary = workbook.create_sheet(SUMMARY_SHEET)

    total_runs = len(rows)
    total_cost = sum(float(r.get("total_cost") or 0) for r in rows)
    total_input_tokens = sum(int(r.get("total_input_tokens") or 0) for r in rows)
    total_output_tokens = sum(int(r.get("total_output_tokens") or 0) for r in rows)
    total_items = sum(int(r.get("total_items") or 0) for r in rows)
    total_revise = sum(int(r.get("revise_count") or 0) for r in rows)
    total_remove = sum(int(r.get("remove_count") or 0) for r in rows)
    total_parse_errors = sum(int(r.get("parse_error_count") or 0) for r in rows)
    avg_cost = total_cost / total_runs if total_runs else 0.0

    summary.append(["Metric", "Value"])
    summary.append(["Total runs", total_runs])
    summary.append(["Total cost (USD)", round(total_cost, 4)])
    summary.append(["Average cost per run (USD)", round(avg_cost, 4)])
    summary.append(["Total input tokens", total_input_tokens])
    summary.append(["Total output tokens", total_output_tokens])
    summary.append(["Total items processed", total_items])
    summary.append(["Total REVISE decisions", total_revise])
    summary.append(["Total REMOVE decisions", total_remove])
    summary.append(["Total PARSE_ERROR items", total_parse_errors])
    summary.append([])
    summary.append(["Per-model breakdown"])
    summary.append(["Model", "Runs", "Total cost (USD)", "Total tokens"])

    by_model: Dict[str, Dict[str, float]] = {}
    for r in rows:
        model = r.get("model_id") or "unknown"
        bucket = by_model.setdefault(model, {"runs": 0, "cost": 0.0, "tokens": 0})
        bucket["runs"] += 1
        bucket["cost"] += float(r.get("total_cost") or 0)
        bucket["tokens"] += int(r.get("total_tokens") or 0)

    for model, stats in sorted(by_model.items()):
        summary.append([model, stats["runs"], round(stats["cost"], 4), stats["tokens"]])

    for column_cells in summary.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        summary.column_dimensions[column_cells[0].column_letter].width = length + 2

    workbook.save(path)
    logger.info("Refreshed '%s' sheet in %s (%d runs)", SUMMARY_SHEET, path, total_runs)
    return str(path)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
    target = sys.argv[1] if len(sys.argv) > 1 else "logs/cost_log.xlsx"
    calculate(target)